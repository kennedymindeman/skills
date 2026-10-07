import sys
from pathlib import Path
from fpdf import FPDF
import docx
from pptx import Presentation

out = Path(sys.argv[1])

def pdf(path, pages):
    p = FPDF()
    p.set_font("Helvetica", size=11)
    for lines in pages:
        p.add_page()
        for line in lines:
            p.multi_cell(0, 6, line)
            p.ln(2)
    p.output(str(path))

pdf(out / "rd-quote-lease-clause" / "lease.pdf", [
    ["RESIDENTIAL LEASE AGREEMENT", "Between Northgate Holdings LLC (Landlord) and Dana Ruiz (Tenant).",
     "1. Premises. Unit 4B, 220 Alder Street, Portland, OR 97209.",
     "2. Notices. All notices go to the addresses listed on the signature page.",
     "3. Rent. Monthly rent is $2,140.00, due on the first day of each month."],
    ["12. Pets. One cat or one dog under 30 lb is permitted with a $400 deposit.",
     "13. Maintenance. Tenant reports needed repairs within 72 hours of discovery.",
     "14. Termination. Either party may terminate this Lease on sixty (60) days' written notice delivered to the address in Section 2, provided that no notice may take effect before 31 March 2027.",
     "15. Governing Law. This Lease is governed by the laws of the State of Oregon."],
])

pdf(out / "rd-behavior-extract-before-visual" / "annual-report.pdf", [
    ["Harbor Point Water Utility - 2025 Annual Report", "Prepared for the Board of Commissioners.",
     "Section 1. Overview. The utility served 48,310 connections in 2025, up from 47,902 in 2024."],
    ["Section 2. Operations.", "Total water produced: 6,812 million gallons.",
     "Billed authorized consumption: 5,627 million gallons.",
     "Non-revenue water for 2025 was 17.4 percent of production, down from 19.1 percent in 2024.",
     "Main breaks repaired: 212."],
    ["Section 3. Capital Program.", "The Eastside reservoir rehabilitation finished in October 2025 at a cost of $3.9 million."],
])

d = docx.Document()
d.add_heading("Services Agreement", 1)
d.add_paragraph("This agreement is made between Larkspur Analytics Inc. (Client) and Fenwick Data Co. (Contractor).")
d.add_paragraph("The Contractor will deliver the data migration described in Exhibit A. Fees and the payment schedule are set out in the table below.")
t = d.add_table(rows=4, cols=2)
for r, (a, b) in enumerate([("Item", "Amount"), ("Phase 1 - discovery", "$38,500.00"),
                             ("Phase 2 - migration", "$145,750.00"), ("Total contract value", "$184,250.00")]):
    t.cell(r, 0).text = a
    t.cell(r, 1).text = b
d.add_paragraph("Invoices are payable net 30.")
d.save(out / "rd-trigger-docx-table" / "services-agreement.docx")

pr = Presentation()
slides = [("Q3 Review", "Platform team - October 2026"),
          ("Incidents", "Two sev-2 outages in August.\nAction: Priya to add alerting on queue depth by Oct 20."),
          ("Roadmap", "Search reindex slipped to Q4.\nAction: Marcus to publish revised reindex plan by Oct 15."),
          ("Hiring", "One SRE offer accepted.\nAction: Lee to schedule onboarding for Nov 3 start.")]
for title, body in slides:
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = body
pr.save(out / "rd-trigger-subagent-brief" / "q3-review.pptx")
